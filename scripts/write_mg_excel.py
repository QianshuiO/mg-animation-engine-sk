#!/usr/bin/env python3
"""Write an MG animation director storyboard Excel file."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


COLUMNS = [
    "\u955c\u53f7",
    "\u65f6\u957f",
    "\u539f\u6587/\u65c1\u767d",
    "RAG\u53c2\u8003",
    "\u753b\u9762\u8bbe\u8ba1",
    "\u955c\u5934\u8fd0\u52a8/\u8f6c\u573a",
    "\u5c4f\u5e55\u6587\u5b57",
    "AI\u63d0\u793a\u8bcd",
]

LEGACY_COLUMNS = ["\u6587\u6848", "\u5206\u955c", "AI\u63d0\u793a\u8bcd"]
COLUMN_WIDTHS = {
    "A": 8,
    "B": 10,
    "C": 34,
    "D": 24,
    "E": 42,
    "F": 30,
    "G": 24,
    "H": 62,
}


def sanitize_filename(text: str, fallback: str = "MG\u52a8\u753b\u811a\u672c") -> str:
    text = re.sub(r"\s+", "", text or "")
    text = re.sub(r'[\\/:*?"<>|]', "", text).strip(" .")
    return (text or fallback)[:20] or fallback


def unique_path(folder: Path, stem: str) -> Path:
    candidate = folder / f"{stem}.xlsx"
    idx = 2
    while candidate.exists():
        candidate = folder / f"{stem}_{idx}.xlsx"
        idx += 1
    return candidate


def load_rows(path: str) -> list[dict[str, str]]:
    if path == "-":
        data = sys.stdin.read()
    else:
        data = Path(path).read_text(encoding="utf-8-sig")
    rows = json.loads(data)
    if not isinstance(rows, list):
        raise SystemExit("rows-json must contain a JSON array.")
    normalized = []
    for idx, item in enumerate(rows, 1):
        if not isinstance(item, dict):
            raise SystemExit("Each row must be a JSON object.")
        row = {
            "\u955c\u53f7": str(item.get("\u955c\u53f7", "") or f"{idx:02d}").strip(),
            "\u65f6\u957f": str(item.get("\u65f6\u957f", "")).strip(),
            "\u539f\u6587/\u65c1\u767d": str(item.get("\u539f\u6587/\u65c1\u767d", "") or item.get("\u6587\u6848", "")).strip(),
            "RAG\u53c2\u8003": str(item.get("RAG\u53c2\u8003", "")).strip(),
            "\u753b\u9762\u8bbe\u8ba1": str(item.get("\u753b\u9762\u8bbe\u8ba1", "") or item.get("\u5206\u955c", "")).strip(),
            "\u955c\u5934\u8fd0\u52a8/\u8f6c\u573a": str(item.get("\u955c\u5934\u8fd0\u52a8/\u8f6c\u573a", "")).strip(),
            "\u5c4f\u5e55\u6587\u5b57": str(item.get("\u5c4f\u5e55\u6587\u5b57", "")).strip(),
            "AI\u63d0\u793a\u8bcd": str(item.get("AI\u63d0\u793a\u8bcd", "")).strip(),
        }
        if any(item.get(col) for col in LEGACY_COLUMNS) and not item.get("RAG\u53c2\u8003"):
            row["RAG\u53c2\u8003"] = "\u57fa\u7840\u63a8\u5bfc"
        if any(item.get(col) for col in LEGACY_COLUMNS) and not item.get("\u65f6\u957f"):
            row["\u65f6\u957f"] = "5s"
        normalized.append(row)
    return normalized


def default_project_dir() -> Path:
    desktop = Path.home() / "Desktop"
    root = desktop if desktop.exists() else Path.home()
    return root / "MG\u52a8\u753b\u811a\u672c\u8f93\u51fa"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project", default="")
    parser.add_argument("--title-source", default="")
    parser.add_argument("--rows-json", required=True)
    args = parser.parse_args()

    project = Path(args.project).expanduser().resolve() if args.project else default_project_dir().resolve()
    output_dir = project / "\u811a\u672c\u5728\u8fd9\u91cc"
    output_dir.mkdir(parents=True, exist_ok=True)

    rows = load_rows(args.rows_json)
    if not rows:
        raise SystemExit("No rows to write.")

    output_path = unique_path(output_dir, sanitize_filename(args.title_source or rows[0].get("\u539f\u6587/\u65c1\u767d", "")))
    wb = Workbook()
    ws = wb.active
    ws.title = "\u5236\u4f5c\u5206\u955c\u811a\u672c"
    ws.append(COLUMNS)
    for row in rows:
        ws.append([row.get(col, "") for col in COLUMNS])

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for idx in range(2, ws.max_row + 1):
        ws.row_dimensions[idx].height = 96
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

    wb.save(output_path)
    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
